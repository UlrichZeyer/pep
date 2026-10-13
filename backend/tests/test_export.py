"""Export (M9): PDF-Aushang, Excel, CSV für die Lohnbuchhaltung."""
import io
import re
from datetime import date, datetime, timezone

import pytest
from openpyxl import load_workbook
from sqlmodel import select

from app.models import Abteilung, Abwesenheit, Einsatz, Mitarbeiter, Plan


@pytest.fixture()
def plan(session):
    """Eigener Testplan KW 42/2026 mit drei Einsätzen und einer Krankmeldung."""
    abt = session.exec(select(Abteilung).order_by(Abteilung.sortierung)).first()
    ma = list(session.exec(select(Mitarbeiter).where(Mitarbeiter.status == "aktiv").order_by(Mitarbeiter.id)))
    p = Plan(name="Export-Test KW 42", von=date(2026, 10, 12), bis=date(2026, 10, 17), status="freigegeben",
             erstellt_am=datetime.now(timezone.utc))  # models.Plan: Default ohne Zeitzone
    session.add(p)
    session.flush()
    einsaetze = [
        Einsatz(plan_id=p.id, mitarbeiter_id=ma[0].id, abteilung_id=abt.id, datum=date(2026, 10, 12),
                beginn="06:00", ende="14:30", pause_min=30),
        Einsatz(plan_id=p.id, mitarbeiter_id=ma[0].id, abteilung_id=abt.id, datum=date(2026, 10, 13),
                beginn="12:15", ende="16:30", pause_min=0, gesperrt=True),
        Einsatz(plan_id=p.id, mitarbeiter_id=ma[1].id, abteilung_id=abt.id, datum=date(2026, 10, 14),
                beginn="07:00", ende="13:00", pause_min=0),
    ]
    krank = Abwesenheit(mitarbeiter_id=ma[2].id, von=date(2026, 10, 15), bis=date(2026, 10, 16), typ="krank",
                        status="genehmigt")
    session.add_all([*einsaetze, krank])
    session.commit()
    yield p, ma, abt
    for e in einsaetze:
        session.delete(e)
    session.delete(krank)
    session.delete(p)
    session.commit()


@pytest.fixture()
def plan_mehrwoechig(session):
    """Plan über Weihnachten/Neujahr (zwei Wochen, Feiertage BW)."""
    abt = session.exec(select(Abteilung)).first()
    ma = session.exec(select(Mitarbeiter).where(Mitarbeiter.status == "aktiv")).first()
    p = Plan(name="Export-Test Weihnachten", von=date(2026, 12, 21), bis=date(2027, 1, 2), erstellt_am=datetime.now(timezone.utc))
    session.add(p)
    session.flush()
    e = Einsatz(plan_id=p.id, mitarbeiter_id=ma.id, abteilung_id=abt.id, datum=date(2026, 12, 23),
                beginn="08:00", ende="17:00", pause_min=45)
    session.add(e)
    session.commit()
    yield p
    session.delete(e)
    session.delete(p)
    session.commit()


def test_pdf(client, plan):
    p, *_ = plan
    r = client.get(f"/api/plaene/{p.id}/export?format=pdf")
    assert r.status_code == 200, r.text
    assert r.content.startswith(b"%PDF")
    assert r.headers["content-type"] == "application/pdf"
    assert r.headers["content-disposition"] == 'attachment; filename="PEP_KW42_2026.pdf"'
    assert len(re.findall(rb"/Type /Page\b", r.content)) == 1


def test_pdf_mehrere_wochen(client, plan_mehrwoechig):
    r = client.get(f"/api/plaene/{plan_mehrwoechig.id}/export?format=pdf")
    assert r.status_code == 200
    assert "PEP_KW52-KW53_2026.pdf" in r.headers["content-disposition"]
    assert len(re.findall(rb"/Type /Page\b", r.content)) >= 2  # je Woche eine Seite


def test_xlsx(client, plan):
    p, ma, abt = plan
    r = client.get(f"/api/plaene/{p.id}/export?format=xlsx")
    assert r.status_code == 200
    assert "PEP_KW42_2026.xlsx" in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Wochenplan", "Einsätze"]
    werte = [str(c.value) for row in wb["Wochenplan"].iter_rows() for c in row if c.value is not None]
    assert f"06:00–14:30 {abt.kuerzel}" in werte
    assert "abwesend" in werte  # Krankheit nur neutral
    assert not any("krank" in w.lower() for w in werte)
    assert wb["Wochenplan"].freeze_panes is not None
    zeilen = list(wb["Einsätze"].iter_rows(values_only=True))
    assert zeilen[0][:3] == ("Persnr", "Name", "Datum")
    assert len(zeilen) == 4
    erste = zeilen[1]
    assert erste[0] == ma[0].persnr and erste[6] == 8.0 and erste[8] == "nein"
    assert zeilen[2][8] == "ja"


def test_xlsx_feiertage_geschlossen(client, plan_mehrwoechig):
    r = client.get(f"/api/plaene/{plan_mehrwoechig.id}/export?format=xlsx")
    ws = load_workbook(io.BytesIO(r.content))["Wochenplan"]
    kopf = [str(c.value) for c in ws[4]]
    assert any(k.startswith("Fr 25.12.") for k in kopf)
    assert not any(k.startswith("So") for k in kopf)  # nur Mo–Sa
    assert "geschlossen" in [c.value for row in ws.iter_rows() for c in row]


def test_csv(client, plan):
    p, ma, abt = plan
    r = client.get(f"/api/plaene/{p.id}/export?format=csv")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/csv")
    assert "PEP_KW42_2026.csv" in r.headers["content-disposition"]
    assert r.content.startswith(b"\xef\xbb\xbf")
    zeilen = r.content.decode("utf-8-sig").splitlines()
    assert zeilen[0] == "persnr;nachname;vorname;datum;beginn;ende;pause_min;netto_std;abteilung;vertragsart"
    assert len(zeilen) == 4
    erste = zeilen[1].split(";")
    assert erste[:3] == [ma[0].persnr, ma[0].nachname, ma[0].vorname]
    assert erste[3:8] == ["2026-10-12", "06:00", "14:30", "30", "8,00"]
    assert erste[8] == abt.name and erste[9]
    assert any(z.split(";")[7] == "4,25" for z in zeilen[1:])


def test_fehler(client, plan):
    p, *_ = plan
    assert client.get("/api/plaene/999999/export?format=pdf").status_code == 404
    r = client.get(f"/api/plaene/{p.id}/export?format=docx")
    assert r.status_code == 422
    assert "Exportformat" in r.json()["detail"]
