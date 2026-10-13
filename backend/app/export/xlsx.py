"""Excel-Arbeitsmappe: Blatt „Wochenplan“ (Matrix) und Blatt „Einsätze“ (Liste)."""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.export.daten import AKZENT, FILIALE, WOCHENTAGE, PlanDaten
from app.util import arbeitsminuten

_HEX = AKZENT.lstrip("#")
KOPF_FUELLUNG = PatternFill("solid", fgColor=_HEX)
KOPF_SCHRIFT = Font(bold=True, color="FFFFFF")
GRUPPE_FUELLUNG = PatternFill("solid", fgColor="F6E3D3")
GESCHLOSSEN_FUELLUNG = PatternFill("solid", fgColor="EFE6DA")
LINIE = Side(style="thin", color="EFE6DA")
RAHMEN = Border(left=LINIE, right=LINIE, top=LINIE, bottom=LINIE)
UNTEN_AKZENT = Border(bottom=Side(style="medium", color=_HEX))


def _wochenplan(ws, d: PlanDaten) -> None:
    p = d.plan
    ws["A1"] = f"Wochenplan – {p.name}"
    ws["A1"].font = Font(bold=True, size=14, color=_HEX)
    ws["A2"] = (f"{FILIALE} · Zeitraum {p.von:%d.%m.%Y} – {p.bis:%d.%m.%Y} · Status: {d.status_text}"
                f" · Stand {d.erstellt:%d.%m.%Y %H:%M}")
    ws["A2"].font = Font(color="7D7268")

    kopf = ["", "Persnr", "Name"] + [
        f"{WOCHENTAGE[t.weekday()]} {t:%d.%m.}" + (f"\n{d.geschlossen[t]}" if t in d.geschlossen else "")
        for t in d.tage] + ["Σ Std."]
    kz = 4
    for i, text in enumerate(kopf, start=1):
        c = ws.cell(row=kz, column=i, value=text)
        c.fill, c.font, c.border = KOPF_FUELLUNG, KOPF_SCHRIFT, RAHMEN
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[kz].height = 30
    letzte = len(kopf)

    r = kz + 1
    for g in d.gruppen:
        ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=g.farbe.lstrip("#"))
        c = ws.cell(row=r, column=2, value=g.name)
        c.font = Font(bold=True, color="2F2620")
        for col in range(2, letzte + 1):
            ws.cell(row=r, column=col).fill = GRUPPE_FUELLUNG
        r += 1
        for z in g.zeilen:
            ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=g.farbe.lstrip("#"))
            ws.cell(row=r, column=2, value=z.mitarbeiter.persnr)
            ws.cell(row=r, column=3, value=z.name)
            max_zeilen = 1
            for i, t in enumerate(d.tage):
                text = d.zelle(z, t)
                c = ws.cell(row=r, column=4 + i, value=text or None)
                c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                if t in d.geschlossen:
                    c.fill = GESCHLOSSEN_FUELLUNG
                    c.font = Font(italic=True, color="7D7268")
                elif text and not z.einsaetze.get(t):
                    c.font = Font(italic=True, color="7A4E2D")
                max_zeilen = max(max_zeilen, text.count("\n") + 1)
            s = ws.cell(row=r, column=letzte, value=round(z.minuten(d.tage) / 60, 2))
            s.number_format = "0.00"
            s.font = Font(bold=True)
            for col in range(2, letzte + 1):
                ws.cell(row=r, column=col).border = RAHMEN
            if max_zeilen > 1:
                ws.row_dimensions[r].height = 15 * max_zeilen
            r += 1

    ws.column_dimensions["A"].width = 1.5
    ws.column_dimensions["B"].width = 8
    ws.column_dimensions["C"].width = 24
    for i in range(len(d.tage)):
        ws.column_dimensions[get_column_letter(4 + i)].width = 17
    ws.column_dimensions[get_column_letter(letzte)].width = 8
    ws.freeze_panes = ws.cell(row=kz + 1, column=4)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.print_title_rows = f"{kz}:{kz}"


def _einsaetze(ws, d: PlanDaten) -> None:
    kopf = ["Persnr", "Name", "Datum", "Beginn", "Ende", "Pause (min)", "Netto-Std.", "Abteilung", "gesperrt",
            "Quelle"]
    ws.append(kopf)
    for c in ws[1]:
        c.fill, c.font = KOPF_FUELLUNG, KOPF_SCHRIFT
        c.alignment = Alignment(horizontal="center")
    for e in d.einsaetze:
        m = d.mitarbeiter[e.mitarbeiter_id]
        ws.append([m.persnr, f"{m.nachname}, {m.vorname}", e.datum, e.beginn, e.ende, e.pause_min,
                   round(arbeitsminuten(e.beginn, e.ende, e.pause_min) / 60, 2), d.abt_name(e.abteilung_id),
                   "ja" if e.gesperrt else "nein", e.quelle])
        r = ws.max_row
        ws.cell(row=r, column=3).number_format = "DD.MM.YYYY"
        ws.cell(row=r, column=7).number_format = "0.00"
    for i, breite in enumerate([8, 24, 12, 8, 8, 11, 11, 20, 9, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = breite
    ws.freeze_panes = "A2"
    if d.einsaetze:
        ws.auto_filter.ref = ws.dimensions


def erzeuge_xlsx(d: PlanDaten) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Wochenplan"
    _wochenplan(ws, d)
    _einsaetze(wb.create_sheet("Einsätze"), d)
    wb.properties.title = f"PEP – {d.plan.name}"
    wb.properties.creator = "PEP"
    puffer = io.BytesIO()
    wb.save(puffer)
    return puffer.getvalue()
